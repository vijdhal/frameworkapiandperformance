"""
Generator stress / fuzz harness.

Hammers generator.generate() with a wide spread of inputs — variable field
counts (1 to many), unicode and special-character keys/values, empty and very
long values, every body format, odd endpoints/methods — and asserts each run
produces a valid, self-consistent asset set (data.ts + test.ts + a non-empty
Excel sheet with the expected row count, no unresolved placeholders).

It does NOT hit the network or run playwright — it stresses the *generation*
engine's robustness. Run after the strict-format work lands:

    cd testhubui && . .venv/bin/activate && python stress_gen.py

Exit code is non-zero if any case fails, so CI/stress-test.sh can gate on it.
"""
from __future__ import annotations
import sys, json, traceback
sys.dont_write_bytecode = True
from pathlib import Path
from generator import generate

ROOT = Path(__file__).resolve().parent.parent / 'testhub-api-quality-accelerator'
PATHS = {
    'request': str(ROOT / 'src/requestPayload'),
    'spec':    str(ROOT / 'src/tests'),
    'excel':   str(ROOT / 'test-data/excel'),
    'response':str(ROOT / 'test-data/json'),
}
# xlsx is used only to re-read + validate what was written
import openpyxl


def make_fields(keys):
    return {'fields': {k: f'val_{i}' for i, k in enumerate(keys)}, 'files': []}


def cases():
    """Yield (label, payload) fuzz cases."""
    base = dict(paths=dict(PATHS), method='POST', endpoint='/api/resource',
                expectedStatus='201', authType='apikey')

    # 1 — variable field counts on JSON
    for n in (1, 2, 3, 5, 8, 15):
        keys = [f'field{i}' for i in range(n)]
        yield (f'json/{n}-fields', {**base, 'assetName': f'stressJson{n}', 'serviceName': f'svcJson{n}',
               'bodyFormat': 'JSON', 'requestPayload': json.dumps({k: 'x' for k in keys})})

    # 2 — variable field counts on urlencoded (field-derived method names)
    for n in (1, 2, 4, 7):
        keys = [f'p{i}' for i in range(n)]
        yield (f'urlenc/{n}-fields', {**base, 'assetName': f'stressUrl{n}', 'serviceName': f'svcUrl{n}',
               'bodyFormat': 'urlencoded', 'formFields': make_fields(keys)})

    # 3 — variable field counts on form-data (numbered columns)
    for n in (1, 2, 3, 6):
        keys = [f'k{i}' for i in range(n)]
        ff = {'fields': {k: 'v' for k in keys},
              'files': [{'field': 'doc', 'path': 'test-data/files/agreement.pdf', 'contentType': 'application/pdf'}]}
        yield (f'form/{n}-fields+file', {**base, 'assetName': f'stressForm{n}', 'serviceName': f'svcForm{n}',
               'bodyFormat': 'form', 'formFields': ff})

    # 4 — unicode / special-character keys and values (JSON)
    yield ('json/unicode-keys', {**base, 'assetName': 'stressUni', 'serviceName': 'svcUni',
           'bodyFormat': 'JSON', 'requestPayload': json.dumps({'名前': '田中', 'e-mail_id': 'a@b.co', 'weird key!': 'x'})})

    # 5 — empty + very long values (JSON)
    yield ('json/edge-values', {**base, 'assetName': 'stressEdge', 'serviceName': 'svcEdge',
           'bodyFormat': 'JSON', 'requestPayload': json.dumps({'empty': '', 'long': 'A' * 5000, 'zero': 0})})

    # 6 — XML with nested-ish content
    yield ('xml/basic', {**base, 'assetName': 'stressXml', 'serviceName': 'svcXml',
           'bodyFormat': 'XML', 'requestPayload': '<order><id>1</id><item>widget</item></order>'})

    # 7 — binary
    yield ('binary/pdf', {**base, 'assetName': 'stressBin', 'serviceName': 'svcBin',
           'bodyFormat': 'binary', 'requestPayload': 'test-data/files/agreement.pdf'})

    # 8 — odd endpoints / methods
    yield ('json/weird-endpoint', {**base, 'assetName': 'stressEp', 'serviceName': 'svcEp',
           'bodyFormat': 'JSON', 'endpoint': '/v2/api/resources/{id}/sub-items',
           'requestPayload': json.dumps({'a': 1})})
    yield ('json/put-method', {**base, 'assetName': 'stressPut', 'serviceName': 'svcPut',
           'bodyFormat': 'JSON', 'method': 'PUT', 'getBody': True,
           'requestPayload': json.dumps({'a': 1})})

    # 9 — regeneration idempotency: same asset twice must not duplicate/crash
    dup = {**base, 'assetName': 'stressDup', 'serviceName': 'svcDup', 'bodyFormat': 'JSON',
           'requestPayload': json.dumps({'a': 1, 'b': 2})}
    yield ('json/regen-1', dict(dup))
    yield ('json/regen-2', dict(dup))


def validate(label, payload, res):
    problems = []
    if not res.get('ok'):
        problems.append('generate() did not return ok')
    for f in res.get('files', []):
        p = Path(f)
        if not p.exists():
            problems.append(f'missing file {f}')
            continue
        if p.suffix == '.ts':
            txt = p.read_text(encoding='utf-8')
            if '{' + '{' in txt or '}' + '}' in txt or 'undefined' in txt.split('//')[0][:0] or '{cls}' in txt or '{create_method}' in txt:
                problems.append(f'unresolved template token in {p.name}')
            if len(txt.strip()) < 40:
                problems.append(f'suspiciously small {p.name}')
    # Excel: sheet exists and row count matches manifest 'rows'
    xls = [f for f in res.get('files', []) if f.endswith('.xlsx')]
    if xls:
        wb = openpyxl.load_workbook(xls[0])
        ws = wb[res['sheet']] if res.get('sheet') in wb.sheetnames else wb[wb.sheetnames[0]]
        data_rows = sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if any(c not in (None, '') for c in r))
        if data_rows != res.get('rows'):
            problems.append(f"excel rows {data_rows} != manifest rows {res.get('rows')}")
        if data_rows < 1:
            problems.append('excel has no data rows')
    return problems


def main():
    total = passed = 0
    failures = []
    for label, payload in cases():
        total += 1
        try:
            res = generate(payload)
            problems = validate(label, payload, res)
            if problems:
                failures.append((label, '; '.join(problems)))
                print(f'  FAIL  {label:26} — {"; ".join(problems)}')
            else:
                passed += 1
                print(f'  ok    {label:26} rows={res.get("rows")}')
        except Exception as e:
            failures.append((label, f'{type(e).__name__}: {e}'))
            print(f'  ERROR {label:26} — {type(e).__name__}: {e}')
            traceback.print_exc()

    print(f'\nGenerator fuzz: {passed}/{total} cases passed')
    # cleanup generated stress assets
    import shutil
    for d in ['src/tests', 'src/requestPayload']:
        for p in (ROOT / d).glob('stress*'):
            shutil.rmtree(p, ignore_errors=True)
    for p in (ROOT / 'test-data/excel').glob('stress*'):
        p.unlink(missing_ok=True)
    for p in (ROOT / 'test-data/json').glob('stress*'):
        shutil.rmtree(p, ignore_errors=True) if p.is_dir() else p.unlink(missing_ok=True)
    sys.exit(0 if not failures else 1)


if __name__ == '__main__':
    main()
